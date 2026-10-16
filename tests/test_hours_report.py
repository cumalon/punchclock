from io import BytesIO

from openpyxl import load_workbook
import pytest

from app import app
from timeclock import database
from timeclock.hours_report import build_hours_report


def _add_punch(employee_id, timestamp, punch_type, *, incident=False, note=""):
    with database.get_connection() as db:
        return db.execute("""
            INSERT INTO punch (employee_id, timestamp, type, method, incident, note)
            VALUES (?, ?, ?, 'pin', ?, ?)
        """, (employee_id, timestamp, punch_type, int(incident), note)).lastrowid


def _add_daily_leave(
    employee_id, leave_type_id, *, status="pending", guaranteed_minutes=30,
    started_at="2026-03-01T11:30:00+01:00",
    ended_at="2026-03-01T12:30:00+01:00",
):
    with database.get_connection() as db:
        decision_id = db.execute("""
            INSERT INTO daily_leave_decision
                (employee_id, leave_type_id, local_date, status, decision_source,
                 total_seconds, approval_policy, auto_approve_minutes,
                 guaranteed_daily_minutes, updated_at)
            VALUES (?, ?, '2026-03-01', ?, 'system', 3600, 'manual', NULL, ?,
                    '2026-03-01T12:30:00+01:00')
        """, (employee_id, leave_type_id, status, guaranteed_minutes)).lastrowid
        return db.execute("""
            INSERT INTO employee_leave
                (employee_id, leave_type_id, started_at, ended_at, status,
                 record_state, daily_decision_id)
            VALUES (?, ?, ?, ?, ?, 'finished', ?)
        """, (
            employee_id, leave_type_id, started_at, ended_at, status, decision_id,
        )).lastrowid


def test_report_includes_both_boundaries_and_separates_work_and_leave(client, employee):
    worker = employee()
    leave_type_id = database.create_leave_type(
        "Esmorzar", guaranteed_daily_minutes=30
    )
    _add_punch(worker["id"], "2026-02-28T23:00:00+01:00", "entrada")
    _add_punch(worker["id"], "2026-02-28T23:30:00+01:00", "sortida")
    _add_punch(worker["id"], "2026-03-01T09:00:00+01:00", "entrada")
    _add_punch(worker["id"], "2026-03-01T10:00:00+01:00", "sortida")
    _add_punch(
        worker["id"], "2026-03-01T10:00:00+01:00", "entrada",
        incident=True, note="Entrada consecutiva",
    )
    _add_punch(worker["id"], "2026-03-01T11:00:00+01:00", "sortida")
    _add_punch(worker["id"], "2026-03-01T11:00:00+01:00", "entrada")
    _add_punch(worker["id"], "2026-03-01T12:00:00+01:00", "sortida")
    _add_daily_leave(worker["id"], leave_type_id)
    _add_punch(worker["id"], "2026-03-02T09:00:00+01:00", "entrada")
    report = build_hours_report("2026-03-01", "2026-03-02")

    summary = report["employees"][0]
    assert summary["worked_seconds"] == 90 * 60
    assert summary["leave_seconds"] == 30 * 60
    assert summary["total_seconds"] == 120 * 60
    assert len(report["days"]) == 2
    assert any(issue["category"] == "Treball pendent de revisió"
               for issue in report["incidents"])
    assert any(issue["category"] == "Excés pendent"
               for issue in report["incidents"])
    assert any(issue["category"] == "Fitxatge incomplet"
               for issue in report["incidents"])
    assert all(day["date"] >= "2026-03-01" for day in report["days"])


def test_report_deduplicates_approved_legacy_and_daily_leave_overlap(client, employee):
    worker = employee()
    leave_type_id = database.create_leave_type(
        "Esmorzar", guaranteed_daily_minutes=30
    )
    _add_daily_leave(worker["id"], leave_type_id)
    with database.get_connection() as db:
        db.execute("""
            INSERT INTO employee_leave
                (employee_id, leave_type_id, started_at, ended_at,
                 record_state, status)
            VALUES (?, ?, '2026-03-01T11:45:00+01:00',
                    '2026-03-01T12:05:00+01:00', 'finished', 'approved')
        """, (worker["id"], leave_type_id))

    report = build_hours_report("2026-03-01", "2026-03-01", worker["id"])
    day = report["days"][0]
    assert day["leave_seconds"] == 35 * 60
    assert day["total_seconds"] == 35 * 60
    assert any(issue["category"] == "Intervals de permís solapats"
               for issue in report["incidents"])


@pytest.mark.parametrize(
    ("leave_kind", "status", "worked_minutes", "leave_minutes"),
    [
        ("daily", "pending", 7 * 60, 30),
        ("daily", "rejected", 7 * 60, 30),
        ("daily", "approved", 7 * 60, 60),
        ("legacy", "pending", 7 * 60, 0),
        ("legacy", "rejected", 7 * 60, 0),
    ],
)
def test_report_subtracts_leave_absence_independently_of_recognition(
    admin_client, client, employee, leave_kind, status,
    worked_minutes, leave_minutes,
):
    worker = employee()
    leave_type_id = database.create_leave_type(
        "Esmorzar", guaranteed_daily_minutes=30
    )
    _add_punch(worker["id"], "2026-03-01T09:00:00+01:00", "entrada")
    _add_punch(worker["id"], "2026-03-01T17:00:00+01:00", "sortida")
    if leave_kind == "daily":
        _add_daily_leave(
            worker["id"], leave_type_id, status=status,
            guaranteed_minutes=30,
        )
    else:
        with database.get_connection() as db:
            db.execute("""
                INSERT INTO employee_leave
                    (employee_id, leave_type_id, started_at, ended_at,
                     record_state, status)
                VALUES (?, ?, '2026-03-01T12:00:00+01:00',
                        '2026-03-01T13:00:00+01:00', 'finished', ?)
            """, (worker["id"], leave_type_id, status))

    params = (
        f"employee_id={worker['id']}&date_from=2026-03-01"
        "&date_to=2026-03-01"
    )
    response = admin_client.get("/api/admin/hours-report?" + params)
    assert response.status_code == 200
    summary = response.get_json()["report"]["employees"][0]
    assert summary["worked_seconds"] == worked_minutes * 60
    assert summary["leave_seconds"] == leave_minutes * 60
    assert summary["total_seconds"] == (worked_minutes + leave_minutes) * 60

    xlsx = admin_client.get("/api/admin/hours-report.xlsx?" + params)
    assert xlsx.status_code == 200
    workbook = load_workbook(BytesIO(xlsx.data), read_only=True)
    row = next(workbook["Resum"].iter_rows(min_row=2, values_only=True))
    assert row[1] == worked_minutes / 60
    assert row[2] == leave_minutes / 60
    assert row[3] == (worked_minutes + leave_minutes) / 60
    daily_row = next(
        workbook["Detall diari"].iter_rows(min_row=2, values_only=True)
    )
    assert daily_row[2] == worked_minutes / 60
    assert daily_row[3] == leave_minutes / 60
    assert daily_row[4] == (worked_minutes + leave_minutes) / 60


def test_approved_daily_decision_recognizes_full_minutes(employee, client):
    worker = employee()
    leave_type_id = database.create_leave_type(
        "Esmorzar", guaranteed_daily_minutes=30
    )
    _add_daily_leave(
        worker["id"], leave_type_id, status="approved", guaranteed_minutes=30
    )

    report = build_hours_report("2026-03-01", "2026-03-01", worker["id"])
    assert report["employees"][0]["leave_seconds"] == 60 * 60
    assert report["employees"][0]["total_seconds"] == 60 * 60


def test_report_flags_cross_midnight_and_unresolved_punches_without_counting(client, employee):
    worker = employee()
    _add_punch(worker["id"], "2026-03-01T23:30:00+01:00", "entrada")
    _add_punch(worker["id"], "2026-03-02T00:30:00+01:00", "sortida")
    leave_type_id = database.create_leave_type("Visita", guaranteed_daily_minutes=60)
    _add_daily_leave(
        worker["id"], leave_type_id,
        started_at="2026-03-01T23:30:00+01:00",
        ended_at="2026-03-02T00:30:00+01:00",
    )

    report = build_hours_report("2026-03-01", "2026-03-02")
    assert report["employees"][0]["worked_seconds"] == 0
    assert report["employees"][0]["leave_seconds"] == 0
    categories = {issue["category"] for issue in report["incidents"]}
    assert "Fitxatge incomplet" in categories
    assert "Permís nocturn" in categories


def test_report_uses_europe_madrid_local_date_for_offsets(client, employee):
    worker = employee()
    _add_punch(worker["id"], "2026-03-01T23:30:00-02:00", "entrada")
    _add_punch(worker["id"], "2026-03-02T03:30:00+01:00", "sortida")

    report = build_hours_report("2026-03-02", "2026-03-02", worker["id"])
    assert report["days"][0]["date"] == "2026-03-02"
    assert report["employees"][0]["worked_seconds"] == 60 * 60


def test_report_selects_late_arrival_by_expected_local_date(client, employee):
    worker = employee()
    leave_type_id = database.create_leave_type("Retard", guaranteed_daily_minutes=0)
    with database.get_connection() as db:
        db.execute("""
            INSERT INTO employee_leave
                (employee_id, leave_type_id, started_at, ended_at, status,
                 record_state, regularization_kind, expected_started_at)
            VALUES (?, ?, '2026-03-02T00:30:00+01:00',
                    '2026-03-02T00:30:00+01:00', 'approved', 'finished',
                    'late_arrival', '2026-03-01T23:30:00+01:00')
        """, (worker["id"], leave_type_id))

    report = build_hours_report("2026-03-01", "2026-03-01", worker["id"])
    assert report["employees"][0]["leave_seconds"] == 0
    assert any(issue["category"] == "Permís nocturn"
               for issue in report["incidents"])


def test_hours_report_endpoints_validate_auth_dates_and_export_xlsx(
    admin_client, client, employee,
):
    worker = employee(name="=1+1")
    _add_punch(worker["id"], "2026-03-01T09:00:00+01:00", "entrada")
    _add_punch(worker["id"], "2026-03-01T10:15:00+01:00", "sortida")

    with app.test_client() as anonymous:
        assert anonymous.get(
            "/api/admin/hours-report?date_from=2026-03-01&date_to=2026-03-01"
        ).status_code == 401
    invalid = admin_client.get(
        "/api/admin/hours-report?date_from=2026-03-02&date_to=2026-03-01"
    )
    assert invalid.status_code == 400
    invalid_employee = admin_client.get(
        "/api/admin/hours-report?employee_id=missing&date_from=2026-03-01&date_to=2026-03-01"
    )
    assert invalid_employee.status_code == 400

    response = admin_client.get(
        "/api/admin/hours-report?employee_id={}&date_from=2026-03-01&date_to=2026-03-01".format(worker["id"])
    )
    assert response.status_code == 200
    report = response.get_json()["report"]
    assert report["employees"][0]["worked_seconds"] == 75 * 60

    exported = admin_client.get(
        "/api/admin/hours-report.xlsx?employee_id={}&date_from=2026-03-01&date_to=2026-03-01".format(worker["id"])
    )
    assert exported.status_code == 200
    assert exported.mimetype == (
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    workbook = load_workbook(BytesIO(exported.data), read_only=True)
    assert workbook.sheetnames == ["Resum", "Detall diari", "Incidències"]
    assert workbook["Resum"]["A2"].value == "'=1+1"
    assert workbook["Resum"]["B2"].value == 1.25
    assert workbook["Detall diari"]["A2"].value == "2026-03-01"

    database.create_admin_user(
        "supervisor", "test-supervisor-password", role="supervisor"
    )
    with app.test_client() as supervisor:
        login = supervisor.post("/api/admin/login", json={
            "username": "supervisor",
            "password": "test-supervisor-password",
        })
        assert login.status_code == 200
        supervisor_report = supervisor.get(
            "/api/admin/hours-report?date_from=2026-03-01&date_to=2026-03-01"
        )
        assert supervisor_report.status_code == 200
